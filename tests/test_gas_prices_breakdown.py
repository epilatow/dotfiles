"""Tests for files/local/libexec/gas-prices-breakdown/gas-prices-breakdown."""

from __future__ import annotations

import contextlib
import gzip
import http.client
import itertools
import json
import re
import subprocess
import sys
import time
import urllib.error
import zipfile
from datetime import date, timedelta
from pathlib import Path
from typing import TYPE_CHECKING, Self
from unittest.mock import patch

import openpyxl  # type: ignore[import-untyped]
import pdfplumber
import pytest
import xlrd  # type: ignore[import-untyped]

REPO_ROOT = Path(__file__).parents[1]
sys.path.insert(0, str(REPO_ROOT / "src"))

import gas_prices_breakdown as gpb  # noqa: E402

if TYPE_CHECKING:
    from collections.abc import Iterator

HELPER = (
    REPO_ROOT
    / "files"
    / "local"
    / "libexec"
    / "gas-prices-breakdown"
    / "gas-prices-breakdown"
)

URL = "https://example.invalid/report.xlsx"


DAY = date(2026, 9, 17)

CODES = sorted(set(gpb.STATE_CODE.values()))


@pytest.fixture(autouse=True)
def no_backoff_sleep() -> Iterator[list[float]]:
    """Record the retry delays instead of waiting them out."""
    slept: list[float] = []
    # The script resolves `time.sleep` at call time, so patching the
    # module here is what its retry loop will find.
    with patch.object(time, "sleep", autospec=True, side_effect=slept.append):
        yield slept


def _http_error(code: int) -> urllib.error.HTTPError:
    return urllib.error.HTTPError(
        URL,
        code,
        "boom",
        {},  # type: ignore[arg-type]
        None,
    )


def _sample_sources() -> tuple[
    gpb.LcfsQuote, gpb.AuctionSettlement, gpb.AuctionSettlement, gpb.OrCfpMonth
]:
    """The priced-program values every emitting test builds against."""
    return (
        gpb.LcfsQuote(date(2026, 9, 16), 84.75, "https://example.invalid"),
        gpb.AuctionSettlement(date(2026, 8, 19), 32.48, "Aug 2026 #48"),
        gpb.AuctionSettlement(date(2026, 9, 2), 39.50, "WA #15"),
        gpb.OrCfpMonth("2026-07", 161.01, "https://example.invalid"),
    )


# ---------------------------------------------------------------------------
# Retry classification
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "exc",
    [
        urllib.error.URLError(OSError(8, "nodename nor servname provided")),
        ConnectionResetError(54, "Connection reset by peer"),
        TimeoutError("timed out"),
        http.client.RemoteDisconnected("closed without response"),
        gzip.BadGzipFile("not a gzipped file"),
        _http_error(500),
        _http_error(503),
        _http_error(429),
        _http_error(408),
    ],
    ids=lambda e: f"{type(e).__name__}{getattr(e, 'code', '')}",
)
def test_transient_failures_are_retried(
    exc: Exception, no_backoff_sleep: list[float]
) -> None:
    """A failure that says nothing about the request gets another go."""
    with patch.object(
        gpb, "_http_bytes_once", autospec=True, side_effect=[exc, b"ok"]
    ) as once:
        assert gpb._http_bytes(URL) == b"ok"
    assert once.call_count == 2
    assert no_backoff_sleep == [gpb.HTTP_BACKOFF_SEC]


@pytest.mark.parametrize("code", [400, 403, 404, 410, 451])
def test_client_errors_are_not_retried(
    code: int, no_backoff_sleep: list[float]
) -> None:
    """A 4xx describes the request; repeating it unchanged is pointless."""
    with (
        patch.object(
            gpb,
            "_http_bytes_once",
            autospec=True,
            side_effect=_http_error(code),
        ) as once,
        pytest.raises(gpb.FetchError) as caught,
    ):
        gpb._http_bytes(URL)
    assert once.call_count == 1
    assert no_backoff_sleep == []
    assert "giving up after" not in str(caught.value)


def test_retries_are_exhausted_then_reported(
    no_backoff_sleep: list[float],
) -> None:
    """A source down for the whole run fails as FetchError, not URLError."""
    exc = urllib.error.URLError(ConnectionResetError(54, "reset by peer"))
    with (
        patch.object(
            gpb, "_http_bytes_once", autospec=True, side_effect=exc
        ) as once,
        pytest.raises(gpb.FetchError) as caught,
    ):
        gpb._http_bytes(URL)
    assert once.call_count == gpb.HTTP_ATTEMPTS
    assert len(no_backoff_sleep) == gpb.HTTP_ATTEMPTS - 1
    message = str(caught.value)
    assert URL in message
    assert "reset by peer" in message
    assert f"giving up after {gpb.HTTP_ATTEMPTS} attempts" in message


def test_backoff_grows_between_attempts(no_backoff_sleep: list[float]) -> None:
    """Successive waits widen, so a slow recovery still gets caught."""
    with (
        patch.object(
            gpb,
            "_http_bytes_once",
            autospec=True,
            side_effect=urllib.error.URLError("down"),
        ),
        pytest.raises(gpb.FetchError),
    ):
        gpb._http_bytes(URL)
    assert no_backoff_sleep == sorted(no_backoff_sleep)
    assert len(set(no_backoff_sleep)) == len(no_backoff_sleep)


def test_first_attempt_success_does_not_wait(
    no_backoff_sleep: list[float],
) -> None:
    """The retry path costs a working run nothing."""
    with patch.object(
        gpb, "_http_bytes_once", autospec=True, return_value=b"body"
    ) as once:
        assert gpb._http_bytes(URL) == b"body"
    assert once.call_count == 1
    assert no_backoff_sleep == []


def test_http_text_decodes_the_retried_body() -> None:
    """The text helper sits on top of the retrying byte helper."""
    with patch.object(
        gpb,
        "_http_bytes_once",
        autospec=True,
        side_effect=[TimeoutError("slow"), b"<html>caf\xc3\xa9</html>"],
    ):
        # Built with chr() so this file stays ASCII; the bytes above
        # are the UTF-8 encoding of the same character.
        e_acute = chr(0xE9)
        assert gpb._http_text(URL) == f"<html>caf{e_acute}</html>"


# ---------------------------------------------------------------------------
# Reporting a persistent failure
# ---------------------------------------------------------------------------


def test_non_xlsx_body_is_reported_not_raised_raw(tmp_path: Path) -> None:
    """A source answering 200 with an error page names itself in the error."""
    not_a_workbook = tmp_path / "report.xlsx"
    not_a_workbook.write_text("<html>503 Service Unavailable</html>")
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb._load_workbook(not_a_workbook, "OR CFP")
    assert "OR CFP" in str(caught.value)
    assert not isinstance(caught.value, zipfile.BadZipFile)


def test_valid_zip_that_is_not_a_workbook_is_reported(tmp_path: Path) -> None:
    """A readable zip carrying the wrong contents is still bad input."""
    not_a_workbook = tmp_path / "report.xlsx"
    with zipfile.ZipFile(not_a_workbook, "w") as archive:
        archive.writestr("readme.txt", "not a spreadsheet")
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb._load_workbook(not_a_workbook, "OR CFP")
    assert "OR CFP" in str(caught.value)


def test_missing_spec_md_exits_one_not_two(tmp_path: Path) -> None:
    """Content that stopped parsing is a source failure, not a usage one."""
    html = tmp_path / "page.html"
    html.write_text("<html>no spec here</html>")
    assert gpb.main(["extract-markdown", str(html)]) == 1


def test_absent_snapshot_to_remove_exits_two(tmp_path: Path) -> None:
    """Naming a snapshot the file lacks is a usage error, not a source one."""
    html = tmp_path / "page.html"
    html.write_text("const SNAPSHOTS = [\n];\n")
    assert gpb.main(["remove-snapshot", str(html), "2026-01-01"]) == 2


def test_fetch_error_is_a_snapshot_error() -> None:
    """One `except` in `main` has to cover both failure kinds."""
    assert issubclass(gpb.FetchError, gpb.SnapshotError)


def test_main_reports_snapshot_error_without_a_traceback(
    tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    """A source failure exits nonzero with one error line."""
    html = tmp_path / "page.html"
    html.write_text("const SNAPSHOTS = [\n];\n")
    with patch.object(
        gpb,
        "cmd_add_snapshot",
        autospec=True,
        side_effect=gpb.FetchError("GET https://example.invalid: down"),
    ):
        exit_code = gpb.main(["add-snapshot", str(html)])
    assert exit_code == 1
    assert [r.getMessage() for r in caplog.records] == [
        "GET https://example.invalid: down"
    ]


def test_unexpected_errors_keep_their_traceback(tmp_path: Path) -> None:
    """Only world conditions are swallowed; a bug still surfaces raw."""
    html = tmp_path / "page.html"
    html.write_text("const SNAPSHOTS = [\n];\n")
    with (
        patch.object(
            gpb,
            "cmd_add_snapshot",
            autospec=True,
            side_effect=ZeroDivisionError("bug"),
        ),
        pytest.raises(ZeroDivisionError),
    ):
        gpb.main(["add-snapshot", str(html)])


def test_unreadable_page_is_reported_not_raised(tmp_path: Path) -> None:
    """The path comes from a job's argv, so a bad one is bad input."""
    missing = tmp_path / "absent.html"
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb.read_html(missing)
    assert str(missing) in str(caught.value)


def test_process_exits_one_with_no_traceback(tmp_path: Path) -> None:
    """End-to-end: the crony job gets an exit code and a readable log.

    Driven through the real entry point so the `sys.exit(main())`
    wiring is covered, and through a failure that needs no network so
    the run costs no retry backoff.
    """
    missing = tmp_path / "absent.html"
    result = subprocess.run(
        [sys.executable, str(HELPER), "list-snapshots", str(missing)],
        capture_output=True,
        text=True,
        check=False,
        timeout=120,
    )
    assert result.returncode == 1
    assert "Traceback" not in result.stderr
    assert "ERROR" in result.stderr
    assert str(missing) in result.stderr


# ---------------------------------------------------------------------------
# Source content that stopped matching
# ---------------------------------------------------------------------------
#
# Each source is scraped out of a page or workbook whose shape the
# publisher can change without notice. A run that cannot find what it
# needs has to say so and leave the snapshot alone, because the
# alternative is writing a plausible-looking number nobody can trace.


@pytest.mark.parametrize(
    "blob",
    [
        b"",
        b"<html>503 Service Unavailable</html>",
        b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600,
    ],
    ids=["empty", "error_page", "corrupt_ole2"],
)
def test_lcfs_damaged_export_is_reported_not_raised(
    tmp_path: Path, blob: bytes
) -> None:
    """xlrd signals damage through several unrelated exception types.

    Only the innermost is an `XLRDError`; a corrupt container and a
    truncated transfer surface as other classes entirely, and any of
    them escaping would be a traceback where the tool promises one
    reported line. The error page is the case that arrives with a 200,
    so nothing before this point has reason to treat it as a failure.
    """
    book = tmp_path / "lcfs.xls"
    book.write_bytes(blob)
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb._read_lcfs_daily(book)
    assert "CA LCFS" in str(caught.value)


def test_lcfs_missing_export_is_reported_not_raised(tmp_path: Path) -> None:
    """A download that never landed is bad input, not a crash."""
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb._read_lcfs_daily(tmp_path / "never-written.xls")
    assert "CA LCFS" in str(caught.value)


AAA_TODAY = date(2026, 9, 17)


def _aaa_page(rows: str) -> str:
    return (
        f"<html>as of {AAA_TODAY.month}/{AAA_TODAY.day}/"
        f"{AAA_TODAY.year % 100:02d} "
        f'<table id="sortable"><tr><th>State</th></tr>{rows}</table>'
    )


def _aaa_row(state: str, price: str) -> str:
    return f'<tr><td><a href="#">{state}</a></td><td>${price}</td></tr>'


def test_aaa_missing_table_is_reported() -> None:
    """The page rendering without its table is a source failure."""
    with (
        patch.object(
            gpb,
            "_http_text",
            autospec=True,
            return_value="<html>as of 9/17/26 maintenance",
        ),
        pytest.raises(gpb.SnapshotError) as caught,
    ):
        gpb.fetch_aaa_retail(AAA_TODAY)
    assert "sortable table not found" in str(caught.value)


def test_aaa_short_table_is_reported() -> None:
    """A partial table would silently drop states from the snapshot."""
    page = _aaa_page(_aaa_row("California", "5.62"))
    with (
        patch.object(gpb, "_http_text", autospec=True, return_value=page),
        pytest.raises(gpb.SnapshotError) as caught,
    ):
        gpb.fetch_aaa_retail(AAA_TODAY)
    assert "expected 51" in str(caught.value)


def test_aaa_unknown_state_name_is_reported() -> None:
    """An unmapped row means the name column changed, not a new state."""
    page = _aaa_page(_aaa_row("Atlantis", "5.62"))
    with (
        patch.object(gpb, "_http_text", autospec=True, return_value=page),
        pytest.raises(gpb.SnapshotError) as caught,
    ):
        gpb.fetch_aaa_retail(AAA_TODAY)
    assert "unrecognized state" in str(caught.value)


def test_aaa_parses_every_mapped_state() -> None:
    """The happy path, at the full width the caller requires."""
    rows = "".join(_aaa_row(name, "3.50") for name in sorted(gpb.STATE_CODE))
    with patch.object(
        gpb, "_http_text", autospec=True, return_value=_aaa_page(rows)
    ):
        retail = gpb.fetch_aaa_retail(AAA_TODAY)
    assert len(retail) == 51
    assert retail["CA"] == 3.50


class _FakeSheet:
    """The two-column shape xlrd exposes for the monitor's export."""

    def __init__(self, rows: list[list[object]]) -> None:
        self._rows = rows
        self.nrows = len(rows)
        self.ncols = max((len(r) for r in rows), default=0)

    def cell_value(self, row: int, col: int) -> object:
        cells = self._rows[row]
        return cells[col] if col < len(cells) else ""


class _FakeBook:
    def __init__(self, rows: list[list[object]]) -> None:
        self._sheet = _FakeSheet(rows)

    def sheet_by_index(self, index: int) -> _FakeSheet:
        assert index == 0
        return self._sheet


@contextlib.contextmanager
def _lcfs_export(rows: list[list[object]]) -> Iterator[None]:
    """Stand in for xlrd reading the downloaded export."""
    with patch.object(
        xlrd, "open_workbook", autospec=True, return_value=_FakeBook(rows)
    ):
        yield


@contextlib.contextmanager
def _lcfs_quotes(quotes: list[tuple[date, float]]) -> Iterator[None]:
    """Stand in for the monitor download and its parsed series."""
    with (
        patch.object(gpb, "_http_bytes", autospec=True, return_value=b"xls"),
        patch.object(
            gpb, "_read_lcfs_daily", autospec=True, return_value=quotes
        ),
    ):
        yield


def test_lcfs_uses_the_newest_quote_on_or_before_the_target(
    tmp_path: Path,
) -> None:
    """The assessment is daily, so the target's own quote is the one."""
    with _lcfs_quotes(
        [
            (date(2026, 9, 14), 84.50),
            (date(2026, 9, 15), 84.65),
            (date(2026, 9, 16), 84.75),
        ]
    ):
        quote = gpb.fetch_lcfs_for_target_date(date(2026, 9, 16), tmp_path)
    assert quote.as_of == date(2026, 9, 16)
    assert quote.usd_per_mt == 84.75


def test_lcfs_does_not_use_a_quote_from_after_the_target(
    tmp_path: Path,
) -> None:
    """Re-running an old snapshot must not reach for a later price."""
    with _lcfs_quotes(
        [(date(2026, 9, 14), 84.50), (date(2026, 9, 16), 84.75)]
    ):
        quote = gpb.fetch_lcfs_for_target_date(date(2026, 9, 15), tmp_path)
    assert quote.as_of == date(2026, 9, 14)


def test_lcfs_tolerates_a_weekend_gap(tmp_path: Path) -> None:
    """Quotes stop on non-trading days; that is not the source going quiet."""
    with _lcfs_quotes([(date(2026, 9, 11), 84.35)]):
        quote = gpb.fetch_lcfs_for_target_date(date(2026, 9, 14), tmp_path)
    assert quote.as_of == date(2026, 9, 11)


def test_lcfs_source_gone_quiet_fails_the_run(tmp_path: Path) -> None:
    """A source that stopped updating cannot describe the target date."""
    target = date(2026, 9, 17)
    stale = target - timedelta(days=gpb.LCFS_STALE_DAYS + 1)
    with (
        _lcfs_quotes([(stale, 71.96)]),
        pytest.raises(gpb.SnapshotError) as caught,
    ):
        gpb.fetch_lcfs_for_target_date(target, tmp_path)
    message = str(caught.value)
    assert "CA LCFS" in message
    assert "days old" in message


def test_lcfs_at_the_freshness_limit_still_answers(tmp_path: Path) -> None:
    """The bound is inclusive, so a run on the limit is not a failure."""
    target = date(2026, 9, 17)
    edge = target - timedelta(days=gpb.LCFS_STALE_DAYS)
    with _lcfs_quotes([(edge, 71.96)]):
        assert gpb.fetch_lcfs_for_target_date(target, tmp_path).as_of == edge


def test_lcfs_target_before_the_series_starts_is_reported(
    tmp_path: Path,
) -> None:
    """The export is a rolling window; older targets are off its front."""
    with (
        _lcfs_quotes([(date(2026, 9, 16), 84.75)]),
        pytest.raises(gpb.SnapshotError) as caught,
    ):
        gpb.fetch_lcfs_for_target_date(date(2026, 1, 5), tmp_path)
    assert "CA LCFS" in str(caught.value)


def test_lcfs_export_rows_are_read_in_date_order(tmp_path: Path) -> None:
    """The caller takes the last usable row, so order is load-bearing."""
    with _lcfs_export(
        [
            ["Date", "California LCFS Carbon Credit (USD/ton)"],
            ["2026-09-16", 84.75],
            ["2026-09-14", 84.50],
            ["2026-09-15", 84.65],
        ]
    ):
        assert gpb._read_lcfs_daily(tmp_path / "lcfs.xls") == [
            (date(2026, 9, 14), 84.50),
            (date(2026, 9, 15), 84.65),
            (date(2026, 9, 16), 84.75),
        ]


@pytest.mark.parametrize(
    "rows",
    [
        [["Date", "Price"]],
        [["Date", "Price"], ["not a date", 84.75]],
        [["Date", "Price"], ["2026-09-16", "n/a"]],
        [["2026-09-16"]],
        [],
    ],
    ids=["header_only", "bad_date", "bad_price", "missing_column", "empty"],
)
def test_lcfs_export_without_usable_rows_is_reported(
    tmp_path: Path, rows: list[list[object]]
) -> None:
    """A reshaped export must not read as an empty price series."""
    with (
        _lcfs_export(rows),
        pytest.raises(gpb.SnapshotError) as caught,
    ):
        gpb._read_lcfs_daily(tmp_path / "lcfs.xls")
    assert "no dated rows" in str(caught.value)


@contextlib.contextmanager
def _or_cfp_sources(
    monthly: dict[tuple[int, int], float],
) -> Iterator[None]:
    """Stand in for the DEQ download and the workbook read."""
    with (
        patch.object(gpb, "_http_bytes", autospec=True, return_value=b"xlsx"),
        patch.object(
            gpb, "_read_or_cfp_monthly", autospec=True, return_value=monthly
        ),
    ):
        yield


def test_or_cfp_picks_the_newest_published_month(tmp_path: Path) -> None:
    """DEQ publishes month M by the end of M+1, so M+2 is when it counts."""
    with _or_cfp_sources(
        {(2026, 5): 150.0, (2026, 6): 155.0, (2026, 7): 161.01}
    ):
        # July becomes available 2026-09-01, so a run the day before
        # still has to settle for June.
        month = gpb.fetch_or_cfp_for_target_date(date(2026, 8, 31), tmp_path)
    assert month.year_month == "2026-06"
    assert month.vwap_per_credit == 155.0


@pytest.mark.parametrize(
    ("year_month", "available_on"),
    [
        ((2026, 1), date(2026, 3, 1)),
        ((2026, 10), date(2026, 12, 1)),
        ((2026, 11), date(2027, 1, 1)),
        ((2026, 12), date(2027, 2, 1)),
    ],
    ids=["jan", "oct", "nov", "dec"],
)
def test_or_cfp_availability_crosses_the_year_boundary(
    tmp_path: Path, year_month: tuple[int, int], available_on: date
) -> None:
    """November and December land in the next year, and must not wrap wrong.

    The month-plus-two arithmetic is the one place a year rollover can
    silently shift a value by twelve months, so each boundary month is
    pinned at the day it becomes usable and the day before.
    """
    year, month = year_month
    with _or_cfp_sources({year_month: 99.0}):
        on_the_day = gpb.fetch_or_cfp_for_target_date(available_on, tmp_path)
    assert on_the_day.year_month == f"{year}-{month:02d}"

    day_before = available_on - timedelta(days=1)
    with (
        _or_cfp_sources({year_month: 99.0}),
        pytest.raises(gpb.SnapshotError) as caught,
    ):
        gpb.fetch_or_cfp_for_target_date(day_before, tmp_path)
    assert "no monthly value available" in str(caught.value)


def test_or_cfp_workbook_with_nothing_published_yet_is_reported(
    tmp_path: Path,
) -> None:
    """A workbook whose months are all too recent cannot answer the date."""
    with (
        _or_cfp_sources({(2026, 7): 161.01}),
        pytest.raises(gpb.SnapshotError) as caught,
    ):
        gpb.fetch_or_cfp_for_target_date(date(2026, 8, 1), tmp_path)
    assert "no monthly value available" in str(caught.value)


# ---------------------------------------------------------------------------
# Hand-maintained auction tables
# ---------------------------------------------------------------------------
#
# These are the only inputs a human has to top up, which makes a missed
# update the likeliest way a wrong number reaches the page: the lookup
# keeps answering with the last entry, and nothing about that looks like
# a failure.

AUCTIONS = [
    (date(2026, 2, 18), 27.94, "Feb 2026 #46"),
    (date(2026, 5, 20), 28.81, "May 2026 #47"),
    (date(2026, 8, 19), 32.48, "Aug 2026 #48"),
]


def test_auction_is_not_used_before_its_results_publish() -> None:
    """A settlement is unknowable until it is certified a week later."""
    lag = gpb.CCA_PUBLICATION_LAG_DAYS
    auction = date(2026, 8, 19)
    day_before = gpb.latest_auction(
        AUCTIONS, auction + timedelta(days=lag - 1), program="T"
    )
    assert day_before.label == "May 2026 #47"
    on_publication = gpb.latest_auction(
        AUCTIONS, auction + timedelta(days=lag), program="T"
    )
    assert on_publication.label == "Aug 2026 #48"


def test_stale_auction_table_fails_the_run() -> None:
    """An un-topped-up table must not keep answering with its last entry."""
    target = date(2026, 8, 19) + timedelta(days=gpb.CCA_STALE_DAYS + 1)
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb.latest_auction(AUCTIONS, target, program="WA CCA")
    message = str(caught.value)
    assert "WA CCA" in message
    assert "days old" in message


def test_auction_table_inside_the_limit_still_answers() -> None:
    """The bound has to tolerate a normal quarterly gap."""
    target = date(2026, 8, 19) + timedelta(days=gpb.CCA_STALE_DAYS)
    assert gpb.latest_auction(AUCTIONS, target, program="T").price == 32.48


def test_target_before_any_published_auction_is_reported() -> None:
    """Nothing to fall back on is a source failure, not a silent zero."""
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb.latest_auction(AUCTIONS, date(2026, 1, 1), program="CARB CCA")
    assert "CARB CCA" in str(caught.value)


@pytest.mark.parametrize("name", sorted(gpb.CCA_AUCTION_TABLES))
def test_shipped_tables_are_ordered_and_unique(name: str) -> None:
    """The lookup walks the table in order and keeps the last match."""
    table = gpb.CCA_AUCTION_TABLES[name]
    dates = [row[0] for row in table]
    assert dates == sorted(dates), f"{name} is out of order"
    assert len(set(dates)) == len(dates), f"{name} has a duplicate date"
    assert all(row[1] > 0 for row in table), f"{name} has a bad price"


def test_emitted_program_name_is_current_and_not_half_renamed() -> None:
    """CARB renamed the program; the page should not lag its own source.

    The name appears in the state block's part label and again in the
    snapshot's notes, written by different functions, so a partial
    rename leaves a snapshot disagreeing with itself. Both the presence
    of the new name and the absence of the old one are asserted, since
    only the pair rules out a half-applied rename.
    """
    lcfs, cca, wa, orm = _sample_sources()
    fragment = gpb.build_ca_nontax(date(2026, 9, 17), lcfs, cca)
    notes = gpb.passthrough_notes(lcfs, cca, wa, orm)
    assert 'n:"Cap-and-Invest pass-through"' in fragment
    assert "Cap-and-Trade pass-through" not in fragment
    assert "CA cap-and-invest" in notes
    assert "cap-and-trade" not in notes
    # CARB still serves its pages under the legacy path, so this one
    # deliberately does not track the rename; a blanket search and
    # replace over the file would break the link.
    assert "cap-and-trade-program/auction-information" in fragment


def test_emitted_ca_nontax_fragment_is_well_formed() -> None:
    """The note is prose inside a quoted JS literal, so it has to escape.

    An unescaped quote or brace here would corrupt the snapshot array
    for every reader of the page, and the damage would not show up in
    any value this script checks.
    """
    lcfs, cca, _wa, _orm = _sample_sources()
    fragment = gpb.build_ca_nontax(date(2026, 9, 17), lcfs, cca)
    assert fragment.count("{") == fragment.count("}")
    assert fragment.count("[") == fragment.count("]")
    # Every double quote must open or close a field, never sit loose
    # inside the prose.
    assert fragment.count('"') % 2 == 0
    for field in ("src:", "note:", "asOf:"):
        assert field in fragment
    note = re.search(r'note:"((?:[^"\\]|\\.)*)"', fragment)
    assert note is not None, "note field is not a parseable string literal"
    assert "Argus" in note.group(1)
    assert "not exact" in note.group(1)
    assert gpb.LCFS_LAST_CARB_WEEK.isoformat() in note.group(1)


# ---------------------------------------------------------------------------
# Freshness of the remaining live sources
# ---------------------------------------------------------------------------


def test_aaa_page_gone_stale_fails_the_run() -> None:
    """Frozen averages look identical to fresh ones without the date."""
    stale = AAA_TODAY - timedelta(days=gpb.AAA_STALE_DAYS + 1)
    page = (
        f"<html>as of {stale.month}/{stale.day}/{stale.year % 100:02d} "
        f'<table id="sortable"><tr><th>State</th></tr></table>'
    )
    with (
        patch.object(gpb, "_http_text", autospec=True, return_value=page),
        pytest.raises(gpb.SnapshotError) as caught,
    ):
        gpb.fetch_aaa_retail(AAA_TODAY)
    message = str(caught.value)
    assert "AAA" in message
    assert str(stale) in message


def test_aaa_page_without_an_as_of_date_is_reported() -> None:
    """The freshness check cannot be skipped just because the stamp moved."""
    rows = "".join(_aaa_row(name, "3.50") for name in sorted(gpb.STATE_CODE))
    page = f'<html><table id="sortable"><tr><th>S</th></tr>{rows}</table>'
    with (
        patch.object(gpb, "_http_text", autospec=True, return_value=page),
        pytest.raises(gpb.SnapshotError) as caught,
    ):
        gpb.fetch_aaa_retail(AAA_TODAY)
    assert "no 'as of' date" in str(caught.value)


def test_aaa_page_dated_within_the_limit_still_answers() -> None:
    """A weekend-old page is normal, not a failure."""
    recent = AAA_TODAY - timedelta(days=gpb.AAA_STALE_DAYS)
    rows = "".join(_aaa_row(name, "3.50") for name in sorted(gpb.STATE_CODE))
    page = (
        f"<html>as of {recent.month}/{recent.day}/{recent.year % 100:02d} "
        f'<table id="sortable"><tr><th>S</th></tr>{rows}</table>'
    )
    with patch.object(gpb, "_http_text", autospec=True, return_value=page):
        assert len(gpb.fetch_aaa_retail(AAA_TODAY)) == 51


def test_or_cfp_source_gone_quiet_fails_the_run(tmp_path: Path) -> None:
    """DEQ publishing nothing new must not read as an unchanged market."""
    with (
        _or_cfp_sources({(2025, 1): 120.0}),
        pytest.raises(gpb.SnapshotError) as caught,
    ):
        gpb.fetch_or_cfp_for_target_date(date(2026, 9, 17), tmp_path)
    message = str(caught.value)
    assert "OR CFP" in message
    assert "days before" in message


def test_or_cfp_normal_publication_lag_is_not_stale(tmp_path: Path) -> None:
    """The bound has to clear the widest gap the +2 month rule produces."""
    # Aug becomes available Oct 1, so a Sep 30 run still has to accept
    # July -- the longest a correctly-published month is ever behind.
    with _or_cfp_sources({(2026, 7): 161.01}):
        month = gpb.fetch_or_cfp_for_target_date(date(2026, 9, 30), tmp_path)
    assert month.year_month == "2026-07"


def test_aaa_page_from_after_the_target_fails_the_run() -> None:
    """Backfilling must not take today's prices for a past date.

    The page serves only current averages, so a run for an old target
    would otherwise stamp a months-old snapshot with this morning's
    numbers while its notes claim they are the target's.
    """
    ahead = AAA_TODAY + timedelta(days=gpb.AAA_STALE_DAYS + 1)
    rows = "".join(_aaa_row(name, "3.50") for name in sorted(gpb.STATE_CODE))
    page = (
        f"<html>as of {ahead.month}/{ahead.day}/{ahead.year % 100:02d} "
        f'<table id="sortable"><tr><th>S</th></tr>{rows}</table>'
    )
    with (
        patch.object(gpb, "_http_text", autospec=True, return_value=page),
        pytest.raises(gpb.SnapshotError) as caught,
    ):
        gpb.fetch_aaa_retail(AAA_TODAY)
    message = str(caught.value)
    assert "AAA" in message
    assert "after" in message


def test_aaa_page_slightly_ahead_of_the_target_is_fine() -> None:
    """A catch-up run days later is normal and must still publish."""
    ahead = AAA_TODAY + timedelta(days=gpb.AAA_STALE_DAYS)
    rows = "".join(_aaa_row(name, "3.50") for name in sorted(gpb.STATE_CODE))
    page = (
        f"<html>as of {ahead.month}/{ahead.day}/{ahead.year % 100:02d} "
        f'<table id="sortable"><tr><th>S</th></tr>{rows}</table>'
    )
    with patch.object(gpb, "_http_text", autospec=True, return_value=page):
        assert len(gpb.fetch_aaa_retail(AAA_TODAY)) == 51


@pytest.mark.parametrize(
    "stamp", ["as of 13/45/26", "as of 2/30/26"], ids=["month", "day"]
)
def test_aaa_unparseable_date_is_reported(stamp: str) -> None:
    """A date-shaped string that is not a date must not reach the check."""
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb._aaa_as_of(f"<html>{stamp} <table id='sortable'></table>")
    assert "unreadable" in str(caught.value)


def test_or_cfp_at_the_freshness_limit_still_answers(tmp_path: Path) -> None:
    """The bound is inclusive, so a run on the limit is not a failure.

    Age is measured from the first of the published month, so the target
    is derived from the month rather than the other way round.
    """
    month = date(2026, 5, 1)
    target = month + timedelta(days=gpb.OR_CFP_STALE_DAYS)
    with _or_cfp_sources({(month.year, month.month): 150.0}):
        got = gpb.fetch_or_cfp_for_target_date(target, tmp_path)
    assert got.year_month == "2026-05"


def test_or_cfp_one_day_past_the_limit_fails(tmp_path: Path) -> None:
    """Pins the constant itself, not merely some value far outside it."""
    month = date(2026, 5, 1)
    target = month + timedelta(days=gpb.OR_CFP_STALE_DAYS + 1)
    with (
        _or_cfp_sources({(month.year, month.month): 150.0}),
        pytest.raises(gpb.SnapshotError) as caught,
    ):
        gpb.fetch_or_cfp_for_target_date(target, tmp_path)
    assert "OR CFP" in str(caught.value)


# ---------------------------------------------------------------------------
# Federal and state motor fuel tax
# ---------------------------------------------------------------------------


def _tax_period(
    period: date = date(2026, 7, 1), rate: float = 0.30
) -> gpb.TaxPeriod:
    return gpb.TaxPeriod(
        period=period,
        per_state=dict.fromkeys(CODES, rate),
        excise=dict.fromkeys(CODES, round(rate - 0.01, 4)),
        federal=0.184,
        source_url="https://example.invalid/fueltaxes.xlsx",
        missing=(),
    )


def _tax_sheet(
    wb: openpyxl.Workbook,
    title: str,
    rates: dict[str, float | None],
    federal: float,
) -> None:
    """Add one half-year sheet in the shape EIA publishes.

    The federal row sits above the state block and carries its total one
    column further right, which is the layout the reader has to cope
    with rather than an artefact of this fixture.
    """
    ws = wb.create_sheet(title)
    ws.append(["Federal and state motor fuel taxes"])
    ws.append(["Federal", 0.183, 0.001, None, federal])
    ws.append([])
    ws.append(["State tax", "Other", "Total State", "State & Federal"])
    for name, code in sorted(gpb.STATE_CODE.items()):
        total = rates.get(code)
        # The excise column is the total less the charges grouped
        # beside it; a blank total leaves the whole row blank.
        excise = None if total is None else round(total - 0.01, 4)
        ws.append([name, excise, None, total])


def _tax_workbook(
    path: Path, sheets: dict[str, dict[str, float | None]], federal: float
) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rates in sheets.items():
        _tax_sheet(wb, title, rates, federal)
    wb.save(path)


def _flat_rates(value: float = 0.30) -> dict[str, float | None]:
    return dict.fromkeys(gpb.STATE_CODE.values(), value)


def _periods(
    tmp_path: Path,
    sheets: dict[str, dict[str, float | None]],
    federal: float = 0.184,
) -> list[gpb.TaxPeriod]:
    book = tmp_path / "fueltaxes.xlsx"
    _tax_workbook(book, sheets, federal)
    return gpb._read_tax_periods(book)


def test_tax_sheets_are_keyed_by_the_date_they_take_effect(
    tmp_path: Path,
) -> None:
    """January and July sheets are the half-years they open."""
    periods = _periods(
        tmp_path, {"July 2026": _flat_rates(), "January 2026": _flat_rates()}
    )
    assert [p.period for p in periods] == [date(2026, 1, 1), date(2026, 7, 1)]


@pytest.mark.parametrize(
    "title",
    [
        "January 2026 (revised)",
        "January 2021_revised",
        "January 2019 ",
    ],
)
def test_tax_sheet_revision_markers_do_not_hide_a_period(
    tmp_path: Path, title: str
) -> None:
    """Editions decorate their own tab names; the period still parses."""
    periods = _periods(tmp_path, {title: _flat_rates()})
    assert periods[0].period.month == 1


def test_tax_state_names_survive_footnotes_and_change_markers(
    tmp_path: Path,
) -> None:
    """EIA brackets footnotes onto names and stars rows that changed."""
    book = tmp_path / "fueltaxes.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("July 2026")
    ws.append(["Federal", 0.183, 0.001, None, 0.184])
    for name, code in sorted(gpb.STATE_CODE.items()):
        decorated = {"CA": f"{name}[4]  ", "OH": f"* {name} "}.get(code, name)
        ws.append([decorated, None, None, 0.30 if code != "CA" else 0.7364])
    wb.save(book)
    period = gpb._read_tax_periods(book)[0]
    assert period.per_state["CA"] == 0.7364
    assert period.per_state["OH"] == 0.30


def test_tax_period_in_force_is_the_newest_one_already_started(
    tmp_path: Path,
) -> None:
    """A rate steps on its effective date and holds until the next."""
    periods = _periods(
        tmp_path,
        {
            "January 2026": _flat_rates(0.5350),
            "July 2026": _flat_rates(0.6450),
        },
    )
    assert gpb.select_tax_period(periods, date(2026, 6, 30)).period == date(
        2026, 1, 1
    )
    assert gpb.select_tax_period(periods, date(2026, 7, 1)).period == date(
        2026, 7, 1
    )


def test_tax_target_before_the_workbook_starts_is_reported(
    tmp_path: Path,
) -> None:
    """A backfill reaching past the published history must not guess."""
    periods = _periods(tmp_path, {"July 2026": _flat_rates()})
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb.select_tax_period(periods, date(2015, 1, 1))
    assert "starts at" in str(caught.value)


def test_tax_table_gone_quiet_fails_the_run(tmp_path: Path) -> None:
    """A table EIA stopped revising answers with its last half-year."""
    periods = _periods(tmp_path, {"January 2026": _flat_rates()})
    target = date(2026, 1, 1) + timedelta(days=gpb.EIA_TAX_STALE_DAYS + 1)
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb.select_tax_period(periods, target)
    assert "stopped publishing" in str(caught.value)


def test_tax_table_at_the_freshness_limit_still_answers(
    tmp_path: Path,
) -> None:
    """Pins the constant itself, not merely some value far inside it."""
    periods = _periods(tmp_path, {"January 2026": _flat_rates()})
    target = date(2026, 1, 1) + timedelta(days=gpb.EIA_TAX_STALE_DAYS)
    assert gpb.select_tax_period(periods, target).period == date(2026, 1, 1)


def test_tax_holiday_blank_does_not_make_other_periods_unreadable(
    tmp_path: Path,
) -> None:
    """One half-year's hole must not cost the other twenty-three."""
    holed = _flat_rates()
    holed["CT"] = None
    periods = _periods(
        tmp_path, {"July 2022": holed, "July 2026": _flat_rates()}
    )
    assert [p.missing for p in periods] == [("CT",), ()]
    assert gpb.select_tax_period(periods, date(2026, 9, 17)).missing == ()


def test_tax_period_missing_a_state_is_refused_when_used(
    tmp_path: Path,
) -> None:
    """A blank reads the same as a holiday, so it is never guessed at."""
    holed = _flat_rates()
    holed["CT"] = None
    periods = _periods(tmp_path, {"July 2022": holed})
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb.select_tax_period(periods, date(2022, 8, 1))
    assert "CT" in str(caught.value)
    assert "suspended" in str(caught.value)


def test_tax_sheet_missing_a_state_row_is_reported(tmp_path: Path) -> None:
    """A dropped row is a reshaped source, not a suspended tax."""
    book = tmp_path / "fueltaxes.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("July 2026")
    ws.append(["Federal", 0.183, 0.001, None, 0.184])
    for name, code in sorted(gpb.STATE_CODE.items()):
        if code != "WY":
            ws.append([name, None, None, 0.30])
    wb.save(book)
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb._read_tax_periods(book)
    assert "no row for WY" in str(caught.value)


def test_tax_sheet_without_a_federal_row_is_reported(
    tmp_path: Path,
) -> None:
    """The federal rate rides the same sheet and is equally required."""
    book = tmp_path / "fueltaxes.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("July 2026")
    for name in sorted(gpb.STATE_CODE):
        ws.append([name, None, None, 0.30])
    wb.save(book)
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb._read_tax_periods(book)
    assert "Federal" in str(caught.value)


def test_tax_workbook_without_dated_sheets_is_reported(
    tmp_path: Path,
) -> None:
    """A reshaped workbook must not read as an empty history."""
    book = tmp_path / "fueltaxes.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Contents")
    wb.save(book)
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb._read_tax_periods(book)
    assert "no dated sheets" in str(caught.value)


def test_tax_workbook_that_is_not_a_workbook_is_reported(
    tmp_path: Path,
) -> None:
    """A source answering 200 with an error page is still bad input."""
    book = tmp_path / "fueltaxes.xlsx"
    book.write_text("<html>503 Service Unavailable</html>")
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb._read_tax_periods(book)
    assert "EIA tax" in str(caught.value)


def test_tax_fetch_downloads_then_selects(tmp_path: Path) -> None:
    """The wrapper the snapshot builder calls, end to end."""
    book = tmp_path / "built.xlsx"
    _tax_workbook(
        book,
        {"January 2026": _flat_rates(0.5350), "July 2026": _flat_rates(0.645)},
        federal=0.184,
    )
    with patch.object(
        gpb, "_http_bytes", autospec=True, return_value=book.read_bytes()
    ):
        got = gpb.fetch_tax_period_for_target_date(date(2026, 9, 17), tmp_path)
    assert got.period == date(2026, 7, 1)
    assert got.per_state["CA"] == 0.645
    assert got.federal == 0.184
    assert got.source_url == gpb.EIA_FUELTAXES_URL


def test_a_blank_excise_cell_is_reported_not_a_key_error() -> None:
    """A blank column must exit 1, not raise from an unattended job."""
    period = _tax_period()
    stripped = dict(period.excise)
    del stripped["IN"]
    holed = gpb.TaxPeriod(
        period=period.period,
        per_state=period.per_state,
        excise=stripped,
        federal=period.federal,
        source_url=period.source_url,
        missing=(),
    )
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb.compose_state_tax(
            "IN",
            DAY,
            holed,
            3.92,
            gpb.InGutRate(date(2026, 9, 1), 0.239, "https://x.invalid"),
            {},
        )
    assert "leaves IN's excise blank" in str(caught.value)


# ---------------------------------------------------------------------------
# Indiana gasoline use tax
# ---------------------------------------------------------------------------


class _FakePage:
    def __init__(self, tables: list[list[list[str | None]]]) -> None:
        self._tables = tables

    def extract_tables(self) -> list[list[list[str | None]]]:
        return self._tables


class _FakePdf:
    """The context-manager-with-pages shape pdfplumber exposes."""

    def __init__(self, tables: list[list[list[str | None]]]) -> None:
        self.pages = [_FakePage(tables)]

    def __enter__(self) -> Self:
        return self

    def __exit__(self, *_exc: object) -> None:
        return None


@contextlib.contextmanager
def _in_gut_notice(tables: list[list[list[str | None]]]) -> Iterator[None]:
    """Stand in for pdfplumber reading the downloaded notice."""
    with patch.object(
        pdfplumber, "open", autospec=True, return_value=_FakePdf(tables)
    ):
        yield


@contextlib.contextmanager
def _in_gut_rates(rates: dict[date, float]) -> Iterator[None]:
    """Stand in for the notice download and its parsed history."""
    with (
        patch.object(gpb, "_http_bytes", autospec=True, return_value=b"pdf"),
        patch.object(
            gpb, "_read_in_gut_rates", autospec=True, return_value=rates
        ),
    ):
        yield


def test_in_gut_reads_the_monthly_history(tmp_path: Path) -> None:
    """The notice carries every month's rate, not only the current one."""
    with _in_gut_notice(
        [
            [
                ["September 1, 2026", "September 30, 2026", "23.9 cents"],
                ["October 1, 2026", "October 31, 2026", "23.8 cents"],
            ]
        ]
    ):
        rates = gpb._read_in_gut_rates(tmp_path / "notice.pdf")
    assert rates == {date(2026, 9, 1): 0.239, date(2026, 10, 1): 0.238}


def test_in_gut_reads_the_legacy_semiannual_table_too(
    tmp_path: Path,
) -> None:
    """Rows before July 2014 run six months and share the layout."""
    with _in_gut_notice(
        [
            [["July 1, 2013", "December 31, 2013", "19.4 cents"]],
            [["October 1, 2026", "October 31, 2026", "23.8 cents"]],
        ]
    ):
        rates = gpb._read_in_gut_rates(tmp_path / "notice.pdf")
    assert rates[date(2013, 7, 1)] == 0.194
    assert rates[date(2026, 10, 1)] == 0.238


def test_in_gut_ignores_rows_that_are_not_rates(tmp_path: Path) -> None:
    """The notice runs prose and headers through the same extractor."""
    with _in_gut_notice(
        [
            [
                ["Period from", "Period to", "Rate Per Gallon"],
                [None, None, None],
                ["Disclaimer: this document is not a statement", "", ""],
                ["October 1, 2026", "October 31, 2026", "23.8 cents"],
            ]
        ]
    ):
        rates = gpb._read_in_gut_rates(tmp_path / "notice.pdf")
    assert rates == {date(2026, 10, 1): 0.238}


def test_in_gut_notice_without_rate_rows_is_reported(
    tmp_path: Path,
) -> None:
    """A reshaped notice must not read as an empty history."""
    with (
        _in_gut_notice([[["Period from", "Period to", "Rate Per Gallon"]]]),
        pytest.raises(gpb.SnapshotError) as caught,
    ):
        gpb._read_in_gut_rates(tmp_path / "notice.pdf")
    assert "no rate rows" in str(caught.value)


@pytest.mark.parametrize(
    ("name", "body"),
    [
        ("error_page", b"<html>503 Service Unavailable</html>"),
        ("empty", b""),
        ("header_only", b"%PDF-1.4\n"),
    ],
)
def test_in_gut_damaged_notice_is_reported_not_raised(
    tmp_path: Path, name: str, body: bytes
) -> None:
    """Every way the download goes wrong exits 1, never a traceback."""
    notice = tmp_path / f"{name}.pdf"
    notice.write_bytes(body)
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb._read_in_gut_rates(notice)
    assert "IN GUT" in str(caught.value)


def test_in_gut_missing_notice_is_reported_not_raised(
    tmp_path: Path,
) -> None:
    """A file that never landed is the same class of failure."""
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb._read_in_gut_rates(tmp_path / "absent.pdf")
    assert "IN GUT" in str(caught.value)


def test_in_gut_uses_the_month_the_snapshot_falls_in(
    tmp_path: Path,
) -> None:
    """The rate steps on the first of the month and holds all month."""
    with _in_gut_rates(
        {
            date(2026, 8, 1): 0.219,
            date(2026, 9, 1): 0.239,
            date(2026, 10, 1): 0.238,
        }
    ):
        got = gpb.fetch_in_gut_for_target_date(date(2026, 9, 17), tmp_path)
    assert got.month_start == date(2026, 9, 1)
    assert got.usd_per_gal == 0.239
    assert got.source_url == gpb.IN_GUT_NOTICE_URL


def test_in_gut_does_not_use_a_month_that_has_not_started(
    tmp_path: Path,
) -> None:
    """October's rate is published in September but is not in force."""
    with _in_gut_rates({date(2026, 9, 1): 0.239, date(2026, 10, 1): 0.238}):
        got = gpb.fetch_in_gut_for_target_date(date(2026, 9, 30), tmp_path)
    assert got.month_start == date(2026, 9, 1)


def test_in_gut_target_before_the_notice_starts_is_reported(
    tmp_path: Path,
) -> None:
    """A backfill reaching past the published history must not guess."""
    with (
        _in_gut_rates({date(2026, 9, 1): 0.239}),
        pytest.raises(gpb.SnapshotError) as caught,
    ):
        gpb.fetch_in_gut_for_target_date(date(1990, 1, 1), tmp_path)
    assert "starts at" in str(caught.value)


def test_in_gut_notice_gone_quiet_fails_the_run(tmp_path: Path) -> None:
    """A notice nobody replaces keeps answering with its last month."""
    month = date(2026, 9, 1)
    target = month + timedelta(days=gpb.IN_GUT_STALE_DAYS + 1)
    with (
        _in_gut_rates({month: 0.239}),
        pytest.raises(gpb.SnapshotError) as caught,
    ):
        gpb.fetch_in_gut_for_target_date(target, tmp_path)
    assert "stopped replacing" in str(caught.value)


def test_in_gut_at_the_freshness_limit_still_answers(tmp_path: Path) -> None:
    """Pins the constant itself, not merely some value far inside it."""
    month = date(2026, 9, 1)
    target = month + timedelta(days=gpb.IN_GUT_STALE_DAYS)
    with _in_gut_rates({month: 0.239}):
        got = gpb.fetch_in_gut_for_target_date(target, tmp_path)
    assert got.month_start == month


# ---------------------------------------------------------------------------
# Indiana tax suspensions
# ---------------------------------------------------------------------------
#
# A published rate table says what a tax is, not whether it is being
# collected: EIA carried Indiana at 64.5 cents for a July 2026 on which
# both its excise and its use tax were suspended by executive order.


def test_notes_name_the_states_whose_rate_is_overridden() -> None:
    """The summary must not claim only Indiana when five are adjusted."""
    period = _tax_period()
    gut = gpb.InGutRate(date(2026, 5, 1), 0.233, "https://x.invalid")
    during = gpb.state_tax_notes(period, gut, date(2026, 5, 21))
    assert "GA" in during and "KY" in during
    after = gpb.state_tax_notes(period, gut, date(2026, 9, 17))
    assert "GA" not in after


def test_indiana_is_not_in_the_flat_override_table() -> None:
    """Its composition path returns before the table is consulted.

    An Indiana entry would be silently ignored, so the invariant the
    TaxAdjustment docstring states is checked rather than trusted.
    """
    assert all(e.state != "IN" for e in gpb.STATE_TAX_ADJUSTMENTS)


def test_overrides_do_not_overlap_within_a_state() -> None:
    """Two windows on one state would make the one reported arbitrary."""
    by_state: dict[str, list[gpb.TaxAdjustment]] = {}
    for entry in gpb.STATE_TAX_ADJUSTMENTS:
        by_state.setdefault(entry.state, []).append(entry)
    for entries in by_state.values():
        entries.sort(key=lambda e: e.start)
        for earlier, later in itertools.pairwise(entries):
            assert earlier.end < later.start


def test_every_override_is_a_reduction_with_an_authority() -> None:
    """A positive figure would be a rate rise, which the table carries."""
    for entry in gpb.STATE_TAX_ADJUSTMENTS:
        assert entry.usd_per_gal < 0
        assert entry.authority
        assert entry.source_url.startswith("https://")
        assert entry.state in gpb.STATE_CODE.values()


@pytest.mark.parametrize(
    ("state", "day", "expected"),
    [
        ("GA", date(2026, 3, 19), None),
        ("GA", date(2026, 3, 20), -0.333),
        ("GA", date(2026, 6, 2), -0.333),
        ("GA", date(2026, 6, 3), None),
        ("KY", date(2026, 5, 11), -0.10),
        ("KY", date(2026, 7, 1), None),
        ("UT", date(2026, 7, 1), -0.06),
        ("IL", date(2026, 7, 1), -0.013),
        ("TX", date(2026, 7, 1), None),
    ],
    ids=lambda v: str(v),
)
def test_override_windows_are_inclusive_at_both_ends(
    state: str, day: date, expected: float | None
) -> None:
    """A window off by a day mis-prices a whole snapshot."""
    got = gpb.adjustment_on(state, day)
    assert (got.usd_per_gal if got else None) == expected


def test_no_override_starts_before_the_sweep_began() -> None:
    """An entry the sweep never covered would be trusted unverified.

    Only the start is bounded. An end may legitimately run past the
    sweep -- Illinois and Utah are fixed by statute through 2026-12-31
    -- because the sweep's range limits which dates get priced, not how
    far a known window may reach.
    """
    for entry in gpb.STATE_TAX_ADJUSTMENTS:
        assert entry.start >= gpb.TAX_ADJUSTMENTS_CHECKED_FROM


def test_a_date_outside_the_sweep_is_refused_in_both_directions() -> None:
    """A date nobody swept and a date with no override look the same."""
    for day in (
        gpb.TAX_ADJUSTMENTS_CHECKED_FROM - timedelta(days=1),
        gpb.TAX_ADJUSTMENTS_CHECKED_THROUGH + timedelta(days=1),
    ):
        with pytest.raises(gpb.SnapshotError) as caught:
            gpb.assert_adjustments_cover(day)
        assert "swept only for" in str(caught.value)


def test_indiana_checked_through_covers_every_shipped_window() -> None:
    """A window reaching past the check would be trusted unverified."""
    ends = [
        end
        for windows in (gpb.IN_EXCISE_SUSPENSIONS, gpb.IN_GUT_SUSPENSIONS)
        for _, end, _ in windows
    ]
    assert max(ends) <= gpb.TAX_ADJUSTMENTS_CHECKED_THROUGH


def test_tax_sheet_carries_the_excise_column_for_indiana(
    tmp_path: Path,
) -> None:
    """Indiana's total is rebuilt from parts, so the narrow column counts."""
    periods = _periods(tmp_path, {"July 2026": _flat_rates(0.645)})
    assert periods[0].excise["IN"] == 0.635
    assert periods[0].per_state["IN"] == 0.645


def _gut(rate: float) -> gpb.InGutRate:
    return gpb.InGutRate(date(2026, 9, 1), rate, "https://example.invalid")


def test_indiana_sums_its_three_components_when_nothing_is_suspended() -> None:
    """Excise plus the flat inspection fee plus that month's use tax."""
    got = gpb.resolve_indiana_tax(date(2026, 2, 19), 0.36, _gut(0.149))
    assert got.total == 0.519
    assert got.excise == 0.36
    assert got.inspection_fee == gpb.IN_OIL_INSPECTION_FEE
    assert got.gut == 0.149
    assert got.suspended == ()


def test_indiana_use_tax_suspension_leaves_the_excise_collected() -> None:
    """The first orders reached the use tax only."""
    got = gpb.resolve_indiana_tax(date(2026, 4, 30), 0.36, _gut(0.172))
    assert got.total == 0.37
    assert got.gut == 0.0
    assert got.excise == 0.36
    assert got.suspended == ("gasoline use tax (EO 26-09 through 26-25-1)",)


def test_indiana_with_both_suspended_keeps_only_the_inspection_fee() -> None:
    """The notice excludes the inspection fee from every suspension."""
    got = gpb.resolve_indiana_tax(date(2026, 9, 17), 0.37, _gut(0.239))
    assert got.total == gpb.IN_OIL_INSPECTION_FEE
    assert got.excise == 0.0
    assert got.gut == 0.0
    assert len(got.suspended) == 2


@pytest.mark.parametrize(
    ("day", "excise_off", "gut_off"),
    [
        (date(2026, 4, 7), False, False),
        (date(2026, 4, 8), False, True),
        (date(2026, 5, 5), False, True),
        (date(2026, 5, 6), True, True),
        (date(2026, 10, 5), True, True),
    ],
    ids=["day_before", "gut_starts", "gut_only", "excise_starts", "last_day"],
)
def test_indiana_suspension_windows_are_inclusive_at_both_ends(
    day: date, excise_off: bool, gut_off: bool
) -> None:
    """A window that is off by a day mis-prices a whole snapshot."""
    got = gpb.resolve_indiana_tax(day, 0.36, _gut(0.20))
    assert (got.excise == 0.0) is excise_off
    assert (got.gut == 0.0) is gut_off


def test_indiana_past_the_checked_window_is_refused() -> None:
    """A lapsed suspension and an extended one look identical here."""
    day = gpb.TAX_ADJUSTMENTS_CHECKED_THROUGH + timedelta(days=1)
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb.assert_adjustments_cover(day)
    assert "swept only for" in str(caught.value)


def test_indiana_on_the_last_checked_day_still_answers() -> None:
    """Pins the constant itself, not merely some value inside it."""
    gpb.assert_adjustments_cover(gpb.TAX_ADJUSTMENTS_CHECKED_THROUGH)
    got = gpb.resolve_indiana_tax(
        gpb.TAX_ADJUSTMENTS_CHECKED_THROUGH, 0.37, _gut(0.239)
    )
    assert got.total == gpb.IN_OIL_INSPECTION_FEE


def test_indiana_suspension_windows_are_ordered_and_disjoint() -> None:
    """Overlapping windows would make the authority reported arbitrary."""
    for windows in (gpb.IN_EXCISE_SUSPENSIONS, gpb.IN_GUT_SUSPENSIONS):
        assert windows == sorted(windows)
        for (_, end, _), (start, _, _) in itertools.pairwise(windows):
            assert end < start


# ---------------------------------------------------------------------------
# Gross-receipts taxes the rate table omits
# ---------------------------------------------------------------------------


def _spot(usd: float, year_month: str = "2026-08") -> gpb.SpotQuote:
    return gpb.SpotQuote(year_month, usd, "EER_TEST")


def _spot_payload(months: dict[str, object]) -> bytes:
    rows = [{"period": p, "value": v} for p, v in months.items()]
    return json.dumps({"response": {"data": rows}}).encode()


def test_hawaii_is_a_share_of_the_pump_price() -> None:
    """The tax lands on gross income, and the pump price is that income."""
    got = gpb.resolve_gross_receipts("HI", DAY, 5.484, None)
    assert got.usd_per_gal == round(gpb.HI_GET_RATE * 5.484, 4)
    assert got.exact is True


def test_connecticut_above_the_cap_is_the_statutory_constant() -> None:
    """Rack sits above spot, so a spot at the cap proves the cap binds."""
    got = gpb.resolve_gross_receipts("CT", DAY, 4.492, _spot(3.213))
    assert got.usd_per_gal == round(gpb.CT_PGET_RATE * gpb.CT_PGET_CAP_USD, 4)
    assert got.exact is True
    assert "cap" in got.basis


def test_connecticut_below_the_cap_is_an_estimate() -> None:
    """Under the cap the benchmark does real work, so it is not exact."""
    got = gpb.resolve_gross_receipts("CT", DAY, 4.492, _spot(2.063, "2026-01"))
    assert got.usd_per_gal == round(gpb.CT_PGET_RATE * 2.063, 4)
    assert got.exact is False
    assert "2026-01" in got.basis


def test_connecticut_exactly_at_the_cap_counts_as_capped() -> None:
    """The boundary belongs to the side where the statute decides."""
    got = gpb.resolve_gross_receipts(
        "CT", DAY, 4.492, _spot(gpb.CT_PGET_CAP_USD)
    )
    assert got.exact is True


@pytest.mark.parametrize(
    ("state", "rate"),
    [("DE", gpb.select_de_hsca_rate(DAY)), ("OH", gpb.OH_PAT_RATE)],
)
def test_wholesale_states_scale_with_the_benchmark(
    state: str, rate: float
) -> None:
    """Neither has a cap, so the benchmark carries the whole figure."""
    got = gpb.resolve_gross_receipts(state, DAY, 4.4, _spot(3.2))
    assert got.usd_per_gal == round(rate * 3.2, 4)
    assert got.exact is False


@pytest.mark.parametrize("state", ["CT", "DE", "OH"])
def test_wholesale_states_without_a_benchmark_are_refused(
    state: str,
) -> None:
    """A missing benchmark must fail, not silently price the tax at zero."""
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb.resolve_gross_receipts(state, DAY, 4.4, None)
    assert state in str(caught.value)


def test_delaware_rate_tracks_the_year_it_was_reset_for() -> None:
    """The rate changes every January 1, so the year decides it."""
    assert gpb.select_de_hsca_rate(date(2025, 6, 1)) == 0.01120
    assert gpb.select_de_hsca_rate(date(2026, 6, 1)) == 0.011902


def test_delaware_rate_picks_the_newest_regardless_of_table_order() -> None:
    """The rate returned and the rate year-checked must be one entry."""
    reversed_table = list(reversed(gpb.DE_HSCA_RATES))
    with patch.object(gpb, "DE_HSCA_RATES", reversed_table):
        assert gpb.select_de_hsca_rate(date(2026, 6, 1)) == 0.011902


def test_delaware_rate_before_the_table_starts_is_reported() -> None:
    """A backfill past the recorded years must not guess."""
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb.select_de_hsca_rate(date(2019, 1, 1))
    assert "table starts at" in str(caught.value)


def test_delaware_rate_for_an_unentered_year_fails_the_run() -> None:
    """A missed January would otherwise price a year at the old rate."""
    newest = max(start for start, _ in gpb.DE_HSCA_RATES)
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb.select_de_hsca_rate(date(newest.year + 1, 6, 1))
    assert "missing a year" in str(caught.value)


def test_delaware_entries_are_one_per_year_effective_new_years_day() -> None:
    """The year comparison is exact only for a table of this shape.

    An entry effective mid-year would make the guard reject every
    earlier date in that year as a missing year, which it would not be.
    """
    starts = [start for start, _ in gpb.DE_HSCA_RATES]
    assert all((s.month, s.day) == (1, 1) for s in starts)
    assert starts == sorted(starts)
    assert len({s.year for s in starts}) == len(starts)


def test_delaware_rate_fails_on_the_first_day_of_an_unentered_year() -> None:
    """The boundary no day count gets right.

    The other hand-maintained figures here are bounded by a count of
    days. That cannot work for a rule stated in calendar years:
    December 31 of a leap year is itself 365 days past its January 1,
    so any count generous enough to accept an entry's own December also
    accepts the next New Year's Day -- the one date on which the new
    rate is certainly in force and certainly missing.
    """
    newest = max(start for start, _ in gpb.DE_HSCA_RATES)
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb.select_de_hsca_rate(date(newest.year + 1, 1, 1))
    assert "missing a year" in str(caught.value)


def test_delaware_rate_on_the_last_day_of_its_year_still_answers() -> None:
    """The far edge of the entry's own year is still the entry's."""
    newest = max(start for start, _ in gpb.DE_HSCA_RATES)
    assert gpb.select_de_hsca_rate(date(newest.year, 12, 31))


def test_spot_uses_the_newest_month_already_published() -> None:
    """A month closes, then lands weeks later; only then is it usable."""
    months = {"2026-06": 3.037, "2026-07": 3.222, "2026-08": 3.213}
    got = gpb.select_spot(months, date(2026, 9, 17), "EER_TEST")
    assert got.year_month == "2026-07"


def test_spot_series_gone_quiet_fails_the_run() -> None:
    """A series nobody updates keeps answering with its last month."""
    months = {"2026-01": 2.063}
    target = date(2026, 1, 1) + timedelta(days=gpb.EIA_SPOT_STALE_DAYS + 1)
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb.select_spot(months, target, "EER_TEST")
    assert "has stopped" in str(caught.value)


def test_spot_at_the_freshness_limit_still_answers() -> None:
    """Pins the constant itself, not merely some value far inside it."""
    months = {"2026-01": 2.063}
    target = date(2026, 1, 1) + timedelta(days=gpb.EIA_SPOT_STALE_DAYS)
    assert gpb.select_spot(months, target, "EER_TEST").year_month == "2026-01"


def test_spot_target_before_the_series_starts_is_reported() -> None:
    """A backfill reaching past the published history must not guess."""
    with pytest.raises(gpb.SnapshotError) as caught:
        gpb.select_spot({"2026-08": 3.2}, date(2020, 1, 1), "EER_TEST")
    assert "starts at" in str(caught.value)


def test_spot_values_arrive_quoted_and_are_still_numbers() -> None:
    """The API quotes its numbers; a string must not read as no data."""
    with patch.object(
        gpb,
        "_http_bytes",
        autospec=True,
        return_value=_spot_payload({"2026-08": "3.213"}),
    ):
        got = gpb.load_spot_series("EER_TEST", "key")
    assert got == {"2026-08": 3.213}


def test_spot_months_without_a_price_are_skipped() -> None:
    """The API sends null for a month it has no price for."""
    with patch.object(
        gpb,
        "_http_bytes",
        autospec=True,
        return_value=_spot_payload({"2026-07": None, "2026-08": "3.213"}),
    ):
        got = gpb.load_spot_series("EER_TEST", "key")
    assert got == {"2026-08": 3.213}


def test_spot_response_that_is_not_the_expected_shape_is_reported() -> None:
    """An error body answering 200 must not read as an empty series."""
    with (
        patch.object(
            gpb, "_http_bytes", autospec=True, return_value=b"<html>502</html>"
        ),
        pytest.raises(gpb.SnapshotError) as caught,
    ):
        gpb.load_spot_series("EER_TEST", "key")
    assert "expected shape" in str(caught.value)


def test_spot_response_with_no_priced_months_is_reported() -> None:
    """A series that returns rows but no prices is still unusable."""
    with (
        patch.object(
            gpb,
            "_http_bytes",
            autospec=True,
            return_value=_spot_payload({"2026-08": None}),
        ),
        pytest.raises(gpb.SnapshotError) as caught,
    ):
        gpb.load_spot_series("EER_TEST", "key")
    assert "no priced months" in str(caught.value)


def test_missing_api_key_names_the_file_and_the_registration(
    tmp_path: Path,
) -> None:
    """A run without the key fails rather than skipping four states."""
    with (
        patch.object(gpb, "EIA_API_KEY_PATH", tmp_path / "absent"),
        pytest.raises(gpb.SnapshotError) as caught,
    ):
        gpb.load_eia_api_key()
    assert "absent" in str(caught.value)
    assert "opendata/register" in str(caught.value)


def test_empty_api_key_is_reported(tmp_path: Path) -> None:
    """An empty file is as unusable as a missing one, and quieter."""
    key = tmp_path / "eia-api-key"
    key.write_text("\n")
    with (
        patch.object(gpb, "EIA_API_KEY_PATH", key),
        pytest.raises(gpb.SnapshotError) as caught,
    ):
        gpb.load_eia_api_key()
    assert "is empty" in str(caught.value)


def test_api_key_is_stripped_of_surrounding_whitespace(
    tmp_path: Path,
) -> None:
    """A key file written by an editor carries a trailing newline."""
    key = tmp_path / "eia-api-key"
    key.write_text("  abc123  \n")
    with patch.object(gpb, "EIA_API_KEY_PATH", key):
        assert gpb.load_eia_api_key() == "abc123"


def test_api_key_never_reaches_the_url(
    caplog: pytest.LogCaptureFixture,
) -> None:
    """The URL reaches the retry log and the error text; the key must not.

    This drives the real transport helper rather than the layer above
    it, because that layer is exactly where a query-string key would
    leak -- mocking it would make this test green against code that
    leaks.
    """
    seen: list[tuple[str, dict[str, str] | None]] = []

    def _fail(url: str, headers: dict[str, str] | None = None) -> bytes:
        seen.append((url, headers))
        raise TimeoutError("timed out")

    with (
        patch.object(
            gpb, "_http_bytes_once", autospec=True, side_effect=_fail
        ),
        caplog.at_level("WARNING"),
        pytest.raises(gpb.FetchError) as caught,
    ):
        gpb.load_spot_series("EER_TEST", "SECRETKEYVALUE")
    assert seen, "the transport helper was never reached"
    for url, headers in seen:
        assert "SECRETKEYVALUE" not in url
        assert (headers or {})["X-Api-Key"] == "SECRETKEYVALUE"
    assert "SECRETKEYVALUE" not in str(caught.value)
    assert "SECRETKEYVALUE" not in caplog.text


# ---------------------------------------------------------------------------
# Whole-file round trips
# ---------------------------------------------------------------------------
#
# Every subcommand that writes shares one path: pull the SNAPSHOTS array
# out of the page, rebuild the array, and put the file back at mode 0644.
# Driving that through main() is what catches a splice that corrupted the
# array or a write that leaked the file's previous mode, neither of which
# any single-function test sees.


def _state_block(code: str, retail: str = "3.500") -> str:
    return (
        f'  {{ code:"{code}", state:"{code}", retail:{retail},\n'
        f'    fixed:{{ total:0.100, parts:[], asOf:"Jul 2025"}},\n'
        f'    adval:{{ pct:0, parts:[], asOf:""}},\n'
        f'    nontax:{{ total:0, parts:[], asOf:""}}}}'
    )


def _snapshot_block(day: str, codes: list[str]) -> str:
    body = ",\n".join(_state_block(c) for c in codes) + ",\n"
    # Shaped like a published snapshot: retail provenance, then a
    # state-tax sentence, then the pass-through section.
    notes = (
        "Fixture snapshot. Retail (col 2) from somewhere. "
        "State taxes (cols 3-4) inherited from a baseline. "
        f"{gpb.PASSTHROUGH_NOTES_MARKER}CA LCFS 2026-09-16 ($84.75/MT)."
    )
    return (
        f'{{\n  date: "{day}",\n'
        "  federalExcise: 0.184,\n"
        '  federalSource: "https://example.invalid/p510",\n'
        '  federalVintage: "Unchanged since 1993.",\n'
        f'  notes: "{notes}",\n'
        f"  data: [\n{body}  ]\n}}"
    )


# ---------------------------------------------------------------------------
# Page string escaping
# ---------------------------------------------------------------------------

# The page is JavaScript source inside a script element, so a value that
# ends a string literal early takes the whole snapshot array with it
# rather than just the row that carried it. These drive the escaper and
# the builders directly; the end-to-end case is with the round trips.


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("plain", "plain"),
        ('a "quoted" word', 'a \\"quoted\\" word'),
        ("back\\slash", "back\\\\slash"),
        ('both \\ and "', 'both \\\\ and \\"'),
        ("two\nlines", "two\\nlines"),
        ("a\ttab", "a\\ttab"),
        ("ends </script> early", "ends \\u003c/script> early"),
        ("opens <!-- a comment", "opens \\u003c!-- a comment"),
        ("breaks\u2028a line", "breaks\\u2028a line"),
        ("splits\u2029a para", "splits\\u2029a para"),
        # The page carries this character elsewhere unescaped, so the
        # escaper must not be the one thing that spells it differently.
        (f"4.0% {gpb.TIMES} retail", f"4.0% {gpb.TIMES} retail"),
    ],
    ids=[
        "plain",
        "quote",
        "backslash",
        "both",
        "newline",
        "tab",
        "script_close",
        "comment_open",
        "line_separator",
        "paragraph_separator",
        "non_ascii_kept",
    ],
)
def test_page_strings_are_escaped_for_javascript(
    raw: str, expected: str
) -> None:
    """The page is JS source; one stray quote breaks the whole array."""
    assert gpb._js_string(raw) == expected


PAGE_URL_CONSTANTS = (
    "LCFS_MONITOR_URL",
    "AAA_URL",
    "CARB_CCA_PROGRAM_URL",
    "WA_CCA_PROGRAM_URL",
    "OR_CFP_PROGRAM_URL",
)


def test_no_single_expression_page_field_skips_the_escaper() -> None:
    """A field added without _js_string is the failure mode here.

    Driving values through main() only covers the fields some source
    happens to feed, so the builders are read directly instead. Only
    fields that are one interpolation and nothing else are checked; a
    field built from several pieces is the end-to-end test's job.

    Adjacent f-string fragments are joined first, because the wrap at
    79 columns routinely splits a long `src:"{...}"` across two of
    them and an unjoined scan would read straight past it.
    """
    builders = {
        chunk.split("(", 1)[0]: chunk
        for chunk in HELPER.read_text().split("\ndef ")
        if chunk.startswith("build_")
    }
    assert set(builders) == {
        "build_ca_nontax",
        "build_or_nontax",
        "build_wa_nontax",
        "build_fixed",
        "build_data_block",
        "build_snapshot_object",
    }, f"the discoverable builders changed: {sorted(builders)}"
    joined = re.sub(r"['\"]\s*\n\s*f?['\"]", "", "\n".join(builders.values()))
    fields = re.findall(r'\w+: ?"\{([^{}]+)\}"', joined)
    assert fields, "the field pattern no longer matches the builders"
    unescaped = [
        expr
        for expr in fields
        if not expr.startswith("_js_string(")
        and not expr.endswith(".isoformat()")
    ]
    assert not unescaped, f"reaching the page unescaped: {unescaped}"


def test_a_quote_in_an_authority_label_does_not_break_the_page() -> None:
    """An authority string is hand-entered, so it can carry anything."""
    tax = gpb.StateTax(
        total=0.30,
        parts=(gpb.TaxPart('He said "no"', 0.30, "https://x.invalid", ""),),
        as_of="2026-07-01",
    )
    rendered = gpb.build_fixed(tax)
    # The unescaped form closes the field early and leaves a bare word
    # where the renderer expects a key; either assertion alone catches
    # a broken escaper.
    assert '\\"no\\"' in rendered
    assert 'n:"He said "no"' not in rendered


def _page(*blocks: str) -> str:
    joined = ",\n".join(blocks)
    return f"<html><script>\nconst SNAPSHOTS = [\n{joined}\n];\n</script>\n"


@contextlib.contextmanager
def _backups_under(home: Path) -> Iterator[None]:
    """Send the pre-write backup to a scratch home, not the real one."""
    with patch.object(Path, "home", autospec=True, return_value=home):
        yield


def test_remove_snapshot_round_trip_leaves_the_survivor_verbatim(
    tmp_path: Path,
) -> None:
    """A rebuilt array must differ from the original in exactly one block."""
    keep = _snapshot_block("2026-09-17", ["CA", "TX"])
    drop = _snapshot_block("2026-09-10", ["CA", "TX"])
    html = tmp_path / "page.html"
    html.write_text(_page(keep, drop))
    with _backups_under(tmp_path):
        assert gpb.main(["remove-snapshot", str(html), "2026-09-10"]) == 0
    after = html.read_text()
    assert gpb.parse_snapshots(after).blocks == [keep]
    assert "2026-09-10" not in after


def test_write_path_forces_the_published_mode(tmp_path: Path) -> None:
    """The page is served over the web; a private mode must not survive."""
    html = tmp_path / "page.html"
    html.write_text(_page(_snapshot_block("2026-09-17", ["CA", "TX"])))
    html.chmod(0o600)
    with _backups_under(tmp_path):
        assert gpb.main(["remove-snapshot", str(html), "2026-09-17"]) == 0
    assert html.stat().st_mode & 0o777 == 0o644


def test_write_path_backs_the_page_up_first(tmp_path: Path) -> None:
    """The backup is the only way back from a bad splice."""
    html = tmp_path / "page.html"
    original = _page(_snapshot_block("2026-09-17", ["CA", "TX"]))
    html.write_text(original)
    state_dir = tmp_path / ".local" / "state" / "gas-prices-breakdown"
    with _backups_under(tmp_path):
        assert gpb.main(["remove-snapshot", str(html), "2026-09-17"]) == 0
    backups = list(state_dir.glob("backup-*.html"))
    assert len(backups) == 1
    assert backups[0].read_text() == original


@contextlib.contextmanager
def _add_snapshot_sources(
    retail: float = 3.5,
    period: gpb.TaxPeriod | None = None,
    gut_month: date = date(2026, 9, 1),
) -> Iterator[None]:
    """Stub every fetch add-snapshot makes, so no test touches a network."""
    lcfs, _cca, _wa, orm = _sample_sources()
    with (
        patch.object(
            gpb,
            "fetch_aaa_retail",
            autospec=True,
            return_value=dict.fromkeys(CODES, retail),
        ),
        patch.object(
            gpb, "fetch_lcfs_for_target_date", autospec=True, return_value=lcfs
        ),
        patch.object(
            gpb,
            "fetch_or_cfp_for_target_date",
            autospec=True,
            return_value=orm,
        ),
        patch.object(
            gpb,
            "fetch_tax_period_for_target_date",
            autospec=True,
            return_value=period or _tax_period(),
        ),
        patch.object(
            gpb,
            "fetch_in_gut_for_target_date",
            autospec=True,
            return_value=gpb.InGutRate(
                gut_month, 0.239, "https://example.invalid/dn02.pdf"
            ),
        ),
        patch.object(
            gpb, "load_eia_api_key", autospec=True, return_value="key"
        ),
        patch.object(
            gpb,
            "load_spot_series",
            autospec=True,
            return_value={
                "2026-02": 2.081,
                "2026-03": 2.952,
                "2026-06": 3.037,
                "2026-07": 3.222,
            },
        ),
    ):
        yield


# The hand-maintained URLs the builders interpolate directly. The many
# others the page carries -- every TaxPart.src, so the per-state tax and
# gross-receipts citations -- reach it through one escaped call site in
# build_fixed, which the authority-label test covers. These are the ones
# with a call site each, so each is driven with a quote of its own.
def test_every_hand_entered_string_reaching_the_page_is_escaped(
    tmp_path: Path,
) -> None:
    """A quote anywhere in the emitted object breaks the whole array.

    The auction labels, the source URLs and the notes sentence are
    hand-entered or composed from hand-entered parts, so each is driven
    with a quote in it rather than trusting that the tax parts standing
    alone prove the rest.
    """
    html = tmp_path / "page.html"
    html.write_text(_page(_snapshot_block(gpb.BASELINE_DATE, CODES)))
    quoted = gpb.AuctionSettlement(
        date(2026, 8, 19), 32.48, 'Aug 2026 "Joint" Auction'
    )
    with (
        _backups_under(tmp_path),
        _add_snapshot_sources(),
        patch.object(
            gpb, "latest_auction", autospec=True, return_value=quoted
        ),
        contextlib.ExitStack() as patched,
    ):
        for name in PAGE_URL_CONSTANTS:
            patched.enter_context(
                patch.object(gpb, name, f'https://x.invalid/"{name}"')
            )
        assert (
            gpb.main(["add-snapshot", str(html), "--target", "2026-09-17"])
            == 0
        )
    built = gpb.parse_snapshots(html.read_text()).blocks[0]
    # Every double quote in the block must be either a field delimiter
    # or escaped; an unescaped one inside a value is what breaks it.
    assert '"Joint"' not in built
    assert '\\"Joint\\"' in built
    for name in PAGE_URL_CONSTANTS:
        assert f'"{name}"' not in built, f"{name} reached the page raw"
        assert f'\\"{name}\\"' in built, f"{name} never reached the page"
    # The multiplication sign reaches the page through both an escaped
    # field and an unescaped one; a page that spells it two ways is a
    # page whose escaper quietly rewrote half of it.
    assert "\\u00d7" not in built
    assert built.count(gpb.TIMES) > 1


def test_a_benchmark_derived_figure_says_so_on_the_page(
    tmp_path: Path,
) -> None:
    """A reader must be able to tell an estimate from the statute.

    Connecticut above its cap and Delaware against a spot benchmark are
    the two shapes, and only one of them is exact.
    """
    html = tmp_path / "page.html"
    html.write_text(_page(_snapshot_block(gpb.BASELINE_DATE, CODES)))
    with _backups_under(tmp_path), _add_snapshot_sources():
        assert (
            gpb.main(["add-snapshot", str(html), "--target", "2026-09-17"])
            == 0
        )
    built = gpb.parse_snapshots(html.read_text()).blocks[0]
    blocks = dict(gpb.split_state_blocks(built))
    assert "Estimated." in blocks["DE"]
    assert "Estimated." not in blocks["CT"]


def test_add_snapshot_round_trip_splices_a_built_block_in(
    tmp_path: Path,
) -> None:
    """The scheduled job's whole path, with only the fetches stubbed."""
    baseline = _snapshot_block(gpb.BASELINE_DATE, CODES)
    html = tmp_path / "page.html"
    html.write_text(_page(baseline))
    with _backups_under(tmp_path), _add_snapshot_sources():
        assert (
            gpb.main(["add-snapshot", str(html), "--target", "2026-09-17"])
            == 0
        )
    blocks = gpb.parse_snapshots(html.read_text()).blocks
    assert len(blocks) == 2
    # Newest first, and the block it inherited from is untouched.
    assert 'date: "2026-09-17"' in blocks[0]
    assert blocks[1] == baseline
    assert gpb.PASSTHROUGH_NOTES_MARKER in blocks[0]
    assert gpb.STATE_TAX_NOTES_MARKER in blocks[0]
    assert "Argus via Neste monitor" in blocks[0]


def test_add_snapshot_drops_the_ad_valorem_field(tmp_path: Path) -> None:
    """State tax reaches the page as one figure, so the percentage goes."""
    html = tmp_path / "page.html"
    html.write_text(_page(_snapshot_block(gpb.BASELINE_DATE, CODES)))
    with _backups_under(tmp_path), _add_snapshot_sources():
        gpb.main(["add-snapshot", str(html), "--target", "2026-09-17"])
    built = gpb.parse_snapshots(html.read_text()).blocks[0]
    assert "adval:" not in built
    assert built.count("fixed:{") == 51
    assert built.count("nontax:{") == 51


def test_add_snapshot_takes_the_state_tax_from_the_table(
    tmp_path: Path,
) -> None:
    """The figure is the published one, not inherited from a baseline."""
    html = tmp_path / "page.html"
    html.write_text(_page(_snapshot_block(gpb.BASELINE_DATE, CODES)))
    with (
        _backups_under(tmp_path),
        _add_snapshot_sources(period=_tax_period(rate=0.4235)),
    ):
        gpb.main(["add-snapshot", str(html), "--target", "2026-09-17"])
    built = gpb.parse_snapshots(html.read_text()).blocks[0]
    texas = next(
        c for code, c in gpb.split_state_blocks(built) if code == "TX"
    )
    assert "fixed:{ total:0.4235," in texas
    # The baseline's own 0.100 must not survive anywhere.
    assert "total:0.1000" not in built


def test_add_snapshot_takes_the_federal_rate_from_the_table(
    tmp_path: Path,
) -> None:
    """The federal figure rides the same sheet as the state ones."""
    html = tmp_path / "page.html"
    html.write_text(_page(_snapshot_block(gpb.BASELINE_DATE, CODES)))
    with _backups_under(tmp_path), _add_snapshot_sources():
        gpb.main(["add-snapshot", str(html), "--target", "2026-09-17"])
    built = gpb.parse_snapshots(html.read_text()).blocks[0]
    assert "federalExcise: 0.1840," in built
    assert 'federalSource: "https://example.invalid/fueltaxes.xlsx"' in built
    assert 'federalVintage: "Federal excise and LUST fee' in built


def test_add_snapshot_composes_indiana_rather_than_taking_the_total(
    tmp_path: Path,
) -> None:
    """Indiana is suspended on this date, so its total is the fee alone."""
    html = tmp_path / "page.html"
    html.write_text(_page(_snapshot_block(gpb.BASELINE_DATE, CODES)))
    with _backups_under(tmp_path), _add_snapshot_sources():
        gpb.main(["add-snapshot", str(html), "--target", "2026-09-17"])
    built = gpb.parse_snapshots(html.read_text()).blocks[0]
    indiana = next(
        c for code, c in gpb.split_state_blocks(built) if code == "IN"
    )
    assert f"fixed:{{ total:{gpb.IN_OIL_INSPECTION_FEE:.4f}," in indiana
    assert "Suspended" in indiana


def test_add_snapshot_adds_back_a_gross_receipts_tax(
    tmp_path: Path,
) -> None:
    """Connecticut's is a sixth of what it charges; omitting it is wrong."""
    html = tmp_path / "page.html"
    html.write_text(_page(_snapshot_block(gpb.BASELINE_DATE, CODES)))
    with _backups_under(tmp_path), _add_snapshot_sources():
        gpb.main(["add-snapshot", str(html), "--target", "2026-09-17"])
    built = gpb.parse_snapshots(html.read_text()).blocks[0]
    conn = next(c for code, c in gpb.split_state_blocks(built) if code == "CT")
    expected = round(0.30 + gpb.CT_PGET_RATE * gpb.CT_PGET_CAP_USD, 4)
    assert f"fixed:{{ total:{expected:.4f}," in conn
    assert "gross earnings" in conn


def test_add_snapshot_applies_a_rate_override(tmp_path: Path) -> None:
    """A suspension the table does not carry must reach the page.

    Georgia's excise holiday is larger than any pass-through this
    script prices, so a snapshot that missed it would be wrong by more
    than every non-tax column put together.
    """
    html = tmp_path / "page.html"
    html.write_text(_page(_snapshot_block(gpb.BASELINE_DATE, CODES)))
    with (
        _backups_under(tmp_path),
        _add_snapshot_sources(
            # Above the 33.3 cent holiday, as Georgia's real rate is.
            period=_tax_period(period=date(2026, 1, 1), rate=0.3405),
            gut_month=date(2026, 5, 1),
        ),
    ):
        assert (
            gpb.main(["add-snapshot", str(html), "--target", "2026-05-21"])
            == 0
        )
    built = gpb.parse_snapshots(html.read_text()).blocks[0]
    georgia = next(
        c for code, c in gpb.split_state_blocks(built) if code == "GA"
    )
    assert f"fixed:{{ total:{round(0.3405 - 0.333, 4):.4f}," in georgia
    assert "HB 1199" in georgia


def test_an_override_larger_than_the_rate_is_reported_by_state(
    tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    """The failure must name whose rate it is, not just a bare number.

    Exit 1 alone proves nothing here: every SnapshotError exits 1, and
    a bad figure used to surface as a miscount of states instead.
    """
    html = tmp_path / "page.html"
    html.write_text(_page(_snapshot_block(gpb.BASELINE_DATE, CODES)))
    with (
        _backups_under(tmp_path),
        # Well under Georgia's 33.3 cent holiday, so its total goes below
        # zero and the check has to say so.
        _add_snapshot_sources(
            period=_tax_period(period=date(2026, 1, 1), rate=0.10),
            gut_month=date(2026, 5, 1),
        ),
        caplog.at_level("ERROR"),
    ):
        assert (
            gpb.main(["add-snapshot", str(html), "--target", "2026-05-21"])
            == 1
        )
    assert "GA on 2026-05-21" in caplog.text
    assert "HB 1199" in caplog.text
    assert "more than the" in caplog.text


def test_add_snapshot_outside_the_swept_range_is_refused(
    tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    """Past the sweep, no override table can say a rate was collected."""
    html = tmp_path / "page.html"
    html.write_text(_page(_snapshot_block(gpb.BASELINE_DATE, CODES)))
    beyond = gpb.TAX_ADJUSTMENTS_CHECKED_THROUGH + timedelta(days=1)
    with (
        _backups_under(tmp_path),
        _add_snapshot_sources(),
        caplog.at_level("ERROR"),
    ):
        assert (
            gpb.main(
                ["add-snapshot", str(html), "--target", beyond.isoformat()]
            )
            == 1
        )
    # Named, so this cannot pass on some other source failing first.
    assert "swept only for" in caplog.text


def test_add_snapshot_dry_run_leaves_the_page_alone(tmp_path: Path) -> None:
    """A dry run builds and sanity-checks but must not touch the file."""
    html = tmp_path / "page.html"
    html.write_text(_page(_snapshot_block(gpb.BASELINE_DATE, CODES)))
    before = html.read_text()
    with _backups_under(tmp_path), _add_snapshot_sources():
        assert (
            gpb.main(
                [
                    "add-snapshot",
                    str(html),
                    "--target",
                    "2026-09-17",
                    "--dry-run",
                ]
            )
            == 0
        )
    assert html.read_text() == before
